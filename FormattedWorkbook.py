import sys

import loguru
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from pandas import DataFrame

from vols_functions import fill_cell_names, adjust_columns_width


class FormattedWorkbook(Workbook):
    def __init__(self, logging_level='ERROR', table_style='TableStyleMedium4'):
        super().__init__()
        self.logging_level = logging_level
        self.logger = loguru.logger
        self.table_style = table_style
        self.excel_cell_names = fill_cell_names()
        self.ws = self.active

    def excel_format_table(self, df: DataFrame, save_sheet_name: str, save_table_name: str):
        """ Метод обеспечивает форматирование листа Excel с таблицей."""
        self.logger.remove()
        self.logger.add(sys.stdout, level=self.logging_level)
        self.logger.info(f'Создаем лист "{save_sheet_name}"')
        self.ws = self.create_sheet(title=f'{save_sheet_name}')
        self.logger.info(f'Заполняем лист "{save_sheet_name}" данными')
        for row in dataframe_to_rows(df, index=False, header=True):
            self.ws.append(row)
        self.logger.info(f'Форматирует таблицу "{save_table_name}"')
        self.logger.debug(f'Таблица для форматирования: A1:{self.excel_cell_names[len(df.columns)]}{len(df) + 1}')
        tab = Table(displayName=f'{save_table_name}',
                    ref=f'A1:{self.excel_cell_names[len(df.columns)]}{len(df) + 1}')
        tab.tableStyleInfo = TableStyleInfo(name=self.table_style, showRowStripes=True, showColumnStripes=True)
        self.logger.info(f'Добавляем таблицу "{save_table_name}" на лист "{save_sheet_name}"')
        self.ws.add_table(tab)
        self.logger.info(f'Выравниваем поля по размеру в таблице "{save_table_name}"')
        self.ws = adjust_columns_width(self.ws)
